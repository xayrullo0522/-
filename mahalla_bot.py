"""
Excel -> veb-sayt avtomatik kirituvchi dastur ("Online Mahalla" va shunga o'xshash saytlar uchun).

Ishlash tartibi:
  1. Brauzer ochiladi, siz saytga o'zingiz kirasiz (login / ERI / OneID).
  2. Dastur Excel faylidagi har bir qatorni o'qiydi.
  3. config.yaml dagi "qadamlar" ro'yxati bo'yicha har bir qator uchun
     formani ochadi, maydonlarni to'ldiradi va saqlaydi.
  4. Har bir qator natijasi Excel'dagi "Holat" ustuniga yoziladi
     (OK / XATO). Qayta ishga tushirsangiz, OK bo'lgan qatorlar o'tkazib yuboriladi.

Foydalanish:
  python mahalla_bot.py                 # asosiy ish (config.yaml bo'yicha)
  python mahalla_bot.py --tanlash       # sahifadagi elementlar selektorini aniqlash rejimi
  python mahalla_bot.py --tekshir       # saqlashdan oldin har safar tasdiq so'raydi
  python mahalla_bot.py --qatorlar 2-10 # faqat 2..10-qatorlarni ishlash
  python mahalla_bot.py --config boshqa.yaml
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import time
from pathlib import Path

try:
    import yaml
    from openpyxl import load_workbook
    from playwright.sync_api import Error as PWError
    from playwright.sync_api import TimeoutError as PWTimeout
    from playwright.sync_api import sync_playwright
except ImportError as e:  # pragma: no cover
    print(f"Kerakli kutubxona topilmadi: {e.name}")
    print("Avval o'rnating:  pip install -r requirements.txt")
    sys.exit(1)


BAZA = Path(__file__).resolve().parent
# Windows konsolida o'zbek harflari/belgilar sabab dastur yiqilmasin
for _oqim in (sys.stdout, sys.stderr):
    try:
        _oqim.reconfigure(errors="replace")
    except AttributeError:
        pass
SHABLON = re.compile(r"\{([^{}]+)\}")


# --------------------------------------------------------------------------- #
#  Yordamchi funksiyalar
# --------------------------------------------------------------------------- #

def chiqar(matn: str = "") -> None:
    print(matn, flush=True)


def config_oqish(yol: Path) -> dict:
    if not yol.exists():
        chiqar(f"Sozlama fayli topilmadi: {yol}")
        sys.exit(1)
    with open(yol, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def qiymat_matnga(qiymat, fmt: str | None = None) -> str:
    """Excel katagidagi qiymatni saytga yoziladigan matnga aylantiradi."""
    if qiymat is None:
        return ""
    if isinstance(qiymat, (dt.datetime, dt.date)):
        return qiymat.strftime(fmt or "%d.%m.%Y")
    if isinstance(qiymat, float) and qiymat.is_integer():
        return str(int(qiymat))
    matn = str(qiymat).strip()
    if fmt and matn:
        # Matn ko'rinishidagi sanani ham formatlashga urinib ko'ramiz
        for kirish in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return dt.datetime.strptime(matn, kirish).strftime(fmt)
            except ValueError:
                pass
    return matn


def shablon_toldirish(matn: str, qator: dict) -> str:
    """'{Ustun nomi}' ko'rinishidagi joylarni qator qiymatlari bilan almashtiradi."""
    def almashtir(m):
        nom = m.group(1)
        if nom not in qator:
            raise KeyError(f"Excel'da '{nom}' ustuni yo'q")
        return qiymat_matnga(qator[nom])
    return SHABLON.sub(almashtir, matn)


def qatorlar_oraligi(matn: str | None):
    if not matn:
        return None
    toplam = set()
    for qism in matn.split(","):
        qism = qism.strip()
        if "-" in qism:
            a, b = qism.split("-", 1)
            toplam.update(range(int(a), int(b) + 1))
        elif qism:
            toplam.add(int(qism))
    return toplam


# --------------------------------------------------------------------------- #
#  Excel bilan ishlash
# --------------------------------------------------------------------------- #

class ExcelJadval:
    def __init__(self, sozlama: dict):
        self.fayl = (BAZA / sozlama.get("fayl", "fuqarolar.xlsx")).resolve()
        if not self.fayl.exists():
            chiqar(f"Excel fayl topilmadi: {self.fayl}")
            sys.exit(1)
        self.natija_fayl = (BAZA / sozlama["natija_fayl"]).resolve() if sozlama.get("natija_fayl") else self.fayl
        self.sarlavha_qatori = int(sozlama.get("sarlavha_qatori", 1))
        self.holat_nomi = sozlama.get("holat_ustuni", "Holat")

        self.wb = load_workbook(self.fayl)
        varaq = sozlama.get("varaq")
        self.ws = self.wb[varaq] if varaq else self.wb.active

        self.ustunlar: dict[str, int] = {}
        for katak in self.ws[self.sarlavha_qatori]:
            if katak.value is not None and str(katak.value).strip():
                self.ustunlar[str(katak.value).strip()] = katak.column

        if self.holat_nomi not in self.ustunlar:
            ustun = self.ws.max_column + 1
            self.ws.cell(row=self.sarlavha_qatori, column=ustun, value=self.holat_nomi)
            self.ustunlar[self.holat_nomi] = ustun

    def qatorlar(self):
        """(qator_raqami, {ustun: qiymat}) juftliklarini qaytaradi; bo'sh qatorlar o'tkaziladi."""
        for r in range(self.sarlavha_qatori + 1, self.ws.max_row + 1):
            qator = {nom: self.ws.cell(row=r, column=c).value for nom, c in self.ustunlar.items()}
            malumot = [v for k, v in qator.items() if k != self.holat_nomi]
            if all(v is None or str(v).strip() == "" for v in malumot):
                continue
            yield r, qator

    def holat_yoz(self, r: int, matn: str) -> None:
        self.ws.cell(row=r, column=self.ustunlar[self.holat_nomi], value=matn)
        self.saqlash()

    def saqlash(self) -> None:
        try:
            self.wb.save(self.natija_fayl)
        except PermissionError:
            chiqar(f"  ! {self.natija_fayl.name} saqlanmadi: fayl Excel'da ochiq bo'lsa, yoping.")


# --------------------------------------------------------------------------- #
#  Brauzer bilan ishlash
# --------------------------------------------------------------------------- #

def brauzer_ochish(pw, sozlama: dict):
    profil = (BAZA / sozlama.get("profil_papkasi", "brauzer_profil")).resolve()
    profil.mkdir(parents=True, exist_ok=True)
    umumiy = dict(
        user_data_dir=str(profil),
        headless=bool(sozlama.get("yashirin", False)),
        viewport=None,
        args=["--start-maximized"],
        slow_mo=int(sozlama.get("sekinlashtirish_ms", 0)),
    )
    kanal = sozlama.get("kanal", "chrome")
    if kanal:
        try:
            return pw.chromium.launch_persistent_context(channel=kanal, **umumiy)
        except PWError:
            chiqar(f"'{kanal}' brauzeri topilmadi, Playwright Chromium ishlatiladi.")
    return pw.chromium.launch_persistent_context(**umumiy)


def joriy_sahifa(context):
    """Eng oxirgi ochilgan (faol) varaqni qaytaradi."""
    sahifalar = [p for p in context.pages if not p.is_closed()]
    return sahifalar[-1] if sahifalar else context.new_page()


class Bajaruvchi:
    """config.yaml dagi qadamlarni bitta Excel qatori uchun bajaradi."""

    def __init__(self, context, sozlama: dict, tekshir: bool):
        self.context = context
        self.kutish_ms = int(float(sozlama.get("kutish_soniya", 15)) * 1000)
        self.qadam_pauza = float(sozlama.get("qadamlar_orasi", 0.3))
        self.tekshir = tekshir

    @property
    def page(self):
        return joriy_sahifa(self.context)

    def element(self, selector: str):
        return self.page.locator(selector).first

    def qiymat_olish(self, qadam: dict, qator: dict) -> str:
        if "ustun" in qadam:
            nom = qadam["ustun"]
            if nom not in qator:
                raise KeyError(f"Excel'da '{nom}' ustuni yo'q")
            qiymat = qiymat_matnga(qator[nom], qadam.get("format"))
        else:
            qiymat = shablon_toldirish(str(qadam.get("qiymat", "")), qator)
        almashtirish = qadam.get("almashtirish")
        if almashtirish:
            qiymat = str(almashtirish.get(qiymat, qiymat))
        return qiymat

    def bajar(self, qadamlar: list[dict], qator: dict) -> None:
        for i, qadam in enumerate(qadamlar, 1):
            amal = qadam.get("amal")
            nomi = qadam.get("izoh") or qadam.get("ustun") or qadam.get("selector") or amal
            try:
                self.bitta_qadam(qadam, qator)
            except Exception as e:
                if qadam.get("majburiy", True):
                    raise RuntimeError(f"{i}-qadam ({amal}: {nomi}) — {e}") from e
                chiqar(f"    - {i}-qadam o'tkazib yuborildi ({nomi}): {str(e).splitlines()[0]}")
            if self.qadam_pauza:
                time.sleep(self.qadam_pauza)

    def variant_bosish(self, variant_sel: str | None, qiymat: str, kut: int) -> None:
        """Ochilgan ro'yxatdan matni qiymatga teng variantni bosadi (avval aniq, keyin qisman moslik).
        Matn selektorga qo'yilmaydi, shuning uchun Navro'z kabi apostrofli qiymatlar ham ishlaydi."""
        nomzodlar = self.page.locator(variant_sel or "li, [role='option'], .option, .dropdown-item")
        aniq = re.compile(r"^\s*" + re.escape(qiymat) + r"\s*$", re.IGNORECASE)
        oxiri = time.time() + kut / 1000
        while True:
            for moslik in (aniq, qiymat):
                v = nomzodlar.filter(has_text=moslik).filter(visible=True)
                if v.count():
                    v.first.click(timeout=kut)
                    return
            if time.time() > oxiri:
                raise RuntimeError(f"ro'yxatda '{qiymat}' varianti topilmadi")
            time.sleep(0.25)

    def bitta_qadam(self, q: dict, qator: dict) -> None:
        amal = q.get("amal")
        kut = int(float(q.get("kutish_soniya", self.kutish_ms / 1000)) * 1000)
        sel = shablon_toldirish(q["selector"], qator) if "selector" in q else None

        # Qiymat bo'sh bo'lsa va "bosh_bolsa_otkaz" yoqilgan bo'lsa — qadam bajarilmaydi
        if amal in ("yozish", "tanlash", "royxat", "fayl") and q.get("bosh_bolsa_otkaz", True):
            if not self.qiymat_olish(q, qator):
                return

        if q.get("saqlash") and self.tekshir:
            javob = input("    Saqlash tugmasini bosaymi? [Enter = ha, n = yo'q]: ").strip().lower()
            if javob in ("n", "no", "yoq", "yo'q"):
                raise RuntimeError("foydalanuvchi saqlashni bekor qildi")

        if amal == "ochish":
            url = shablon_toldirish(q["url"], qator)
            self.page.goto(url, timeout=max(kut, 30000), wait_until="domcontentloaded")

        elif amal == "bosish":
            el = self.element(sel)
            el.wait_for(state="visible", timeout=kut)
            el.scroll_into_view_if_needed(timeout=kut)
            el.click(timeout=kut, force=bool(q.get("majburan")))

        elif amal == "yozish":
            qiymat = self.qiymat_olish(q, qator)
            el = self.element(sel)
            el.wait_for(state="visible", timeout=kut)
            el.click(timeout=kut)
            if q.get("sekin", False):
                # Maskali (masalan sana, telefon, JSHSHIR) maydonlar uchun harfma-harf yozish
                el.press("Control+A")
                el.press("Delete")
                el.press_sequentially(qiymat, delay=int(q.get("harf_orasi_ms", 60)))
            else:
                el.fill(qiymat, timeout=kut)
            if q.get("enter"):
                el.press("Enter")
            if q.get("tab"):
                el.press("Tab")

        elif amal == "tanlash":
            # Oddiy HTML <select>
            qiymat = self.qiymat_olish(q, qator)
            el = self.element(sel)
            el.wait_for(state="attached", timeout=kut)
            try:
                el.select_option(label=qiymat, timeout=kut)
            except PWError:
                el.select_option(value=qiymat, timeout=kut)

        elif amal == "royxat":
            # Maxsus (div/ul asosidagi) ochiluvchi ro'yxat: bosib ochish, kerak bo'lsa qidirish, variantni bosish
            qiymat = self.qiymat_olish(q, qator)
            el = self.element(sel)
            el.wait_for(state="visible", timeout=kut)
            el.click(timeout=kut)
            if q.get("qidiruv_selector"):
                qidiruv = self.element(q["qidiruv_selector"])
                qidiruv.fill(qiymat, timeout=kut)
            elif q.get("yozib_qidirish"):
                self.page.keyboard.type(qiymat, delay=40)
            time.sleep(float(q.get("variant_kutish", 0.5)))
            self.variant_bosish(q.get("variant_selector"), qiymat, kut)

        elif amal == "belgilash":
            # Checkbox / radio. 'ustun' berilsa, qiymat "ha/1/true/x" bo'lsa belgilanadi
            el = self.element(sel)
            el.wait_for(state="attached", timeout=kut)
            if "ustun" in q or "qiymat" in q:
                kerak = self.qiymat_olish(q, qator).lower() in ("1", "ha", "true", "x", "+", "yes", "да", "bor")
            else:
                kerak = True
            el.set_checked(kerak, timeout=kut, force=True)

        elif amal == "fayl":
            yol = Path(self.qiymat_olish(q, qator))
            if not yol.is_absolute():
                yol = BAZA / yol
            self.element(sel).set_input_files(str(yol), timeout=kut)

        elif amal == "tugma":
            self.page.keyboard.press(q["tugma"])

        elif amal == "kutish":
            if sel:
                holat = q.get("holat", "visible")
                self.element(sel).wait_for(state=holat, timeout=kut)
            elif q.get("url"):
                self.page.wait_for_url(re.compile(q["url"]), timeout=kut)
            else:
                self.page.wait_for_load_state(q.get("holat", "networkidle"), timeout=kut)

        elif amal == "pauza":
            time.sleep(float(q.get("soniya", 1)))

        elif amal == "tekshirish":
            # Sahifada xato xabari chiqqan-chiqmaganini tekshirish
            if sel and self.page.locator(sel).first.is_visible():
                matn = self.page.locator(sel).first.inner_text().strip()
                raise RuntimeError(f"saytdagi xabar: {matn[:200]}")

        elif amal == "qolda":
            # Masalan, captcha yoki ERI tasdig'i uchun
            input(f"    {q.get('xabar', 'Qolda bajaring va Enter bosing')} ")

        else:
            raise ValueError(f"noma'lum amal: {amal!r}")


# --------------------------------------------------------------------------- #
#  Selektor aniqlash rejimi
# --------------------------------------------------------------------------- #

TANLASH_JS = r"""
(() => {
  if (window.__selTanlash) return;
  window.__selTanlash = true;
  const esc = s => (window.CSS && CSS.escape) ? CSS.escape(s) : s.replace(/([^\w-])/g, '\\$1');
  const unik = (sel) => { try { return document.querySelectorAll(sel).length === 1; } catch (e) { return false; } };
  function yol(el) {
    const tag = el.tagName.toLowerCase();
    if (el.id && !/\d{3,}/.test(el.id) && unik('#' + esc(el.id))) return '#' + esc(el.id);
    for (const a of ['name', 'data-testid', 'formcontrolname', 'placeholder', 'aria-label', 'type']) {
      const v = el.getAttribute(a);
      if (v) { const s = `${tag}[${a}="${v.replace(/"/g, '\\"')}"]`; if (unik(s)) return s; }
    }
    const qism = [];
    let e = el;
    while (e && e.nodeType === 1 && e !== document.body) {
      let s = e.tagName.toLowerCase();
      if (e.id && !/\d{3,}/.test(e.id)) { qism.unshift('#' + esc(e.id)); break; }
      const ota = e.parentElement;
      if (ota) {
        const bir = [...ota.children].filter(c => c.tagName === e.tagName);
        if (bir.length > 1) s += `:nth-of-type(${bir.indexOf(e) + 1})`;
      }
      qism.unshift(s);
      e = ota;
    }
    return qism.join(' > ');
  }
  document.addEventListener('click', (ev) => {
    if (!ev.altKey) return;
    ev.preventDefault(); ev.stopPropagation();
    const el = ev.target;
    const matn = (el.innerText || el.value || '').trim().replace(/\s+/g, ' ').slice(0, 60);
    const tag = el.tagName.toLowerCase();
    const old = el.style.outline;
    el.style.outline = '3px solid red';
    setTimeout(() => el.style.outline = old, 1200);
    window.__selTanlandi({
      css: yol(el), tag, matn,
      turi: el.getAttribute('type') || '',
      url: location.href,
    });
  }, true);
})();
"""


def yaml_matn(s: str) -> str:
    """Matnni config.yaml ga to'g'ridan-to'g'ri qo'yish mumkin bo'lgan ko'rinishda qaytaradi."""
    return "'" + s.replace("'", "''") + "'"


def tanlash_rejimi(context, boshlash_url: str) -> None:
    def tanlandi(_source, info):
        tavsiya = {"input": "yozish", "textarea": "yozish", "select": "tanlash"}.get(info["tag"], "bosish")
        if info["turi"] in ("checkbox", "radio"):
            tavsiya = "belgilash"
        chiqar("-" * 60)
        chiqar(f"Sahifa : {info['url']}")
        chiqar(f"Element: <{info['tag']}> {info['matn']!r}")
        chiqar(f"  - amal: {tavsiya}")
        chiqar(f"    selector: {yaml_matn(info['css'])}")
        if info["matn"] and info["tag"] in ("button", "a", "span", "li", "div"):
            chiqar(f"    # yoki matn bo'yicha:  selector: {yaml_matn('text=' + info['matn'])}")

    context.expose_binding("__selTanlandi", tanlandi)
    context.add_init_script(TANLASH_JS)
    page = joriy_sahifa(context)
    if boshlash_url:
        page.goto(boshlash_url, wait_until="domcontentloaded")
    for p in context.pages:
        try:
            p.evaluate(TANLASH_JS)
        except PWError:
            pass

    chiqar("=" * 60)
    chiqar("SELEKTOR ANIQLASH REJIMI")
    chiqar("  1) Saytga kiring, kerakli forma sahifasini oching.")
    chiqar("  2) Klaviaturada ALT tugmasini bosib turib, kerakli maydonga sichqoncha bilan bosing.")
    chiqar("  3) Bu yerda config.yaml ga qo'yish uchun tayyor qator chiqadi.")
    chiqar("  Tugatish uchun brauzerni yoping yoki bu oynada Ctrl+C bosing.")
    chiqar("=" * 60)
    try:
        while context.pages:
            context.pages[0].wait_for_timeout(300)
    except (KeyboardInterrupt, PWError):
        pass


# --------------------------------------------------------------------------- #
#  Asosiy jarayon
# --------------------------------------------------------------------------- #

def asosiy() -> None:
    ap = argparse.ArgumentParser(description="Excel ma'lumotlarini saytga avtomatik kiritish")
    ap.add_argument("--config", default="config.yaml", help="sozlama fayli (standart: config.yaml)")
    ap.add_argument("--tanlash", action="store_true", help="selektor aniqlash rejimi")
    ap.add_argument("--tekshir", action="store_true", help="saqlashdan oldin tasdiq so'rash")
    ap.add_argument("--qatorlar", help="faqat shu Excel qatorlari, masalan: 2-10,15")
    ap.add_argument("--hammasi", action="store_true", help="OK bo'lgan qatorlarni ham qayta kiritish")
    args = ap.parse_args()

    cfg = config_oqish((BAZA / args.config).resolve())
    sayt = cfg.get("sayt", {})
    brauzer_cfg = cfg.get("brauzer", {})

    with sync_playwright() as pw:
        context = brauzer_ochish(pw, brauzer_cfg)

        if args.tanlash:
            tanlash_rejimi(context, sayt.get("kirish_url", ""))
            return

        jadval = ExcelJadval(cfg.get("excel", {}))
        qadamlar = cfg.get("qadamlar") or []
        if not qadamlar:
            chiqar("config.yaml da 'qadamlar' bo'limi bo'sh.")
            return

        page = joriy_sahifa(context)
        if sayt.get("kirish_url"):
            page.goto(sayt["kirish_url"], wait_until="domcontentloaded")
        chiqar("=" * 60)
        chiqar("Brauzerda saytga kiring (login / ERI / OneID).")
        input("Kirib bo'lgach, shu oynada Enter bosing... ")

        tanlangan = qatorlar_oraligi(args.qatorlar)
        qator_pauza = float(brauzer_cfg.get("qatorlar_orasi", 1))
        xato_skrin = BAZA / "xatolar"
        ok = xato = otkaz = 0

        barcha = list(jadval.qatorlar())
        chiqar(f"Excel'da {len(barcha)} ta qator topildi. Boshlandi...")
        bajaruvchi = Bajaruvchi(context, brauzer_cfg, args.tekshir)

        try:
            for n, (r, qator) in enumerate(barcha, 1):
                if tanlangan and r not in tanlangan:
                    continue
                holat = str(qator.get(jadval.holat_nomi) or "")
                if holat.startswith("OK") and not args.hammasi:
                    otkaz += 1
                    continue

                nom = " ".join(qiymat_matnga(qator.get(u)) for u in cfg.get("excel", {}).get("korsatish_ustunlari", []))
                chiqar(f"[{n}/{len(barcha)}] {r}-qator {nom}".rstrip())
                try:
                    bajaruvchi.bajar(qadamlar, qator)
                    jadval.holat_yoz(r, f"OK {dt.datetime.now():%Y-%m-%d %H:%M}")
                    ok += 1
                    chiqar("    ✔ kiritildi")
                except Exception as e:
                    xato += 1
                    xabar = str(e).splitlines()[0][:300]
                    jadval.holat_yoz(r, f"XATO: {xabar}")
                    chiqar(f"    ✖ {xabar}")
                    try:
                        xato_skrin.mkdir(exist_ok=True)
                        bajaruvchi.page.screenshot(path=str(xato_skrin / f"qator_{r}.png"))
                    except Exception:
                        pass
                    if brauzer_cfg.get("xatoda_toxtash", False):
                        input("    Xato. Davom etish uchun Enter bosing (to'xtatish: Ctrl+C)... ")
                time.sleep(qator_pauza)
        except KeyboardInterrupt:
            chiqar("\nFoydalanuvchi to'xtatdi.")
        finally:
            jadval.saqlash()
            chiqar("=" * 60)
            chiqar(f"Tayyor. Muvaffaqiyatli: {ok}, xato: {xato}, avval kiritilgan (o'tkazildi): {otkaz}")
            chiqar(f"Natija: {jadval.natija_fayl}")
            if xato:
                chiqar(f"Xato qatorlar rasmlari: {xato_skrin}")
            if not brauzer_cfg.get("yashirin"):
                try:
                    input("Brauzerni yopish uchun Enter bosing... ")
                except (KeyboardInterrupt, EOFError):
                    pass
            context.close()


if __name__ == "__main__":
    asosiy()
